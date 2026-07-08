import os
import re


class TextExtractor:
    """Extract plain text from PDF, DOCX, and TXT files."""

    def extract(self, filepath: str) -> str:
        ext = os.path.splitext(filepath)[1].lower()
        if ext == '.pdf':
            return self._extract_pdf(filepath)
        elif ext in ('.docx', '.doc'):
            return self._extract_docx(filepath)
        elif ext == '.txt':
            return self._extract_txt(filepath)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

    def _extract_pdf(self, filepath: str) -> str:
        try:
            import pdfplumber
            text_parts = []
            with pdfplumber.open(filepath) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
            text = "\n".join(text_parts)
            return self._clean(text)
        except Exception:
            # Fallback to PyPDF2
            import PyPDF2
            text_parts = []
            with open(filepath, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
            return self._clean("\n".join(text_parts))

    def _extract_docx(self, filepath: str) -> str:
        from docx import Document
        doc = Document(filepath)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return self._clean("\n".join(paragraphs))

    def _extract_txt(self, filepath: str) -> str:
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                return self._clean(f.read())
        except UnicodeDecodeError:
            import chardet
            with open(filepath, 'rb') as f:
                raw = f.read()
            detected = chardet.detect(raw)
            encoding = detected.get('encoding', 'latin-1') or 'latin-1'
            return self._clean(raw.decode(encoding, errors='replace'))

    def _clean(self, text: str) -> str:
        # Remove excessive whitespace and special characters
        text = re.sub(r'\r\n', '\n', text)
        text = re.sub(r'\r', '\n', text)
        text = re.sub(r'\n{3,}', '\n\n', text)
        text = re.sub(r'[ \t]{2,}', ' ', text)
        text = re.sub(r'[^\x00-\x7F\u00C0-\u024F\u0400-\u04FF]+', ' ', text)
        return text.strip()
